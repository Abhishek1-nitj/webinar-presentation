from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side


SRC = Path("/Users/abhishekkumar/Downloads/2.Data Connection Masterclass.xlsm")
OUT = Path("/Users/abhishekkumar/Downloads/2.Data Connection Masterclass - Consolidation Macro.xlsm")


HEADERS = [
    "Order ID", "Customer ID", "Product ID", "Order Date", "Qty",
    "Product Name", "Category", "Unit Price", "Supplier",
    "Customer Name", "City", "Segment", "Revenue",
]


def rows_by_key(ws, key_col=1):
    headers = [c.value for c in ws[1]]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[key_col - 1]:
            out[row[key_col - 1]] = dict(zip(headers, row))
    return out


def build():
    wb = load_workbook(SRC, keep_vba=True)
    for name in ("Connected_Data",):
        if name in wb.sheetnames:
            del wb[name]

    orders = wb["Orders"]
    products = rows_by_key(wb["Products"])
    customers = rows_by_key(wb["Customers"])
    out = wb.create_sheet("Connected_Data", 3)

    out.append(HEADERS)
    for src_row in orders.iter_rows(min_row=2, values_only=True):
        order_id, customer_id, product_id, order_date, qty = src_row[:5]
        product = products.get(product_id, {})
        customer = customers.get(customer_id, {})
        unit_price = product.get("Unit Price")
        revenue = (qty or 0) * (unit_price or 0)
        out.append([
            order_id, customer_id, product_id, order_date, qty,
            product.get("Product Name"), product.get("Category"), unit_price, product.get("Supplier"),
            customer.get("Customer Name"), customer.get("City"), customer.get("Segment"), revenue,
        ])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in out[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    widths = [14, 14, 13, 13, 8, 24, 18, 12, 18, 22, 16, 14, 13]
    for i, width in enumerate(widths, 1):
        out.column_dimensions[out.cell(1, i).column_letter].width = width

    for row in out.iter_rows(min_row=2, max_row=out.max_row, max_col=len(HEADERS)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")
        row[4].number_format = "#,##0"
        row[7].number_format = "#,##0"
        row[12].number_format = "#,##0"

    out.freeze_panes = "A2"
    out.auto_filter.ref = out.dimensions

    q = wb["Questions"]
    q["A1"] = "Build Connected Data"
    q["A2"] = "Macro source is saved beside the project: output/data_connection_consolidation.bas"
    q["A3"] = "This workbook copy already contains the generated Connected_Data tab."
    q["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    q["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    q["A1"].alignment = Alignment(horizontal="center")
    q.column_dimensions["A"].width = 80

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)


def verify():
    wb = load_workbook(OUT, data_only=False, keep_vba=True)
    ws = wb["Connected_Data"]
    assert ws.max_row == 101, ws.max_row
    assert [ws.cell(1, i).value for i in range(1, 14)] == HEADERS

    products = rows_by_key(wb["Products"])
    customers = rows_by_key(wb["Customers"])
    for r in (2, 10, 50, 101):
        product_id = ws.cell(r, 3).value
        customer_id = ws.cell(r, 2).value
        qty = ws.cell(r, 5).value
        unit_price = products[product_id]["Unit Price"]
        assert ws.cell(r, 6).value == products[product_id]["Product Name"]
        assert ws.cell(r, 10).value == customers[customer_id]["Customer Name"]
        assert ws.cell(r, 13).value == qty * unit_price
    print(f"Verified {ws.max_row - 1} connected rows in {OUT}")


if __name__ == "__main__":
    build()
    verify()
